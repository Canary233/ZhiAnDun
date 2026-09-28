# -*- coding: utf-8 -*-
"""
报告模板引导与占位符解析
========================
- 启动时把随源码分发的内置标准 Word 模板（与原工具一致的 {{占位符}} 模板）
  幂等地登记到 templates 表并设为默认，做到开箱即用。
- 解析任意 .docx 模板中使用到的 {{占位符}}，供界面回显，让用户上传模板后
  立即看到“哪些内容会被自动填充”。
"""
import os
import re
import shutil

import database as db
from config import TEMPLATE_DIR, BUILTIN_TEMPLATE_DIR
from core.helpers import now_display

# 内置标准模板（随源码分发，不放在运行时 data 目录）
BUILTIN_NAME = "标准模板（与原工具一致）"
BUILTIN_FILE_NAME = "builtin_standard.docx"
BUILTIN_SOURCE = os.path.join(str(BUILTIN_TEMPLATE_DIR), "标准漏洞报告模板.docx")

# 系统支持的全部占位符（中文键，与 core/exporter.py 的 context 对应）
SUPPORTED_PLACEHOLDERS = [
    "单位名称", "漏洞名称", "系统名称", "域名", "漏洞地址", "漏洞描述",
    "检测过程", "漏洞等级", "漏洞危害", "整改建议", "归属证明", "报告日期", "单位落款",
]

_PH_RE = re.compile(r"\{\{\s*([^{}%]+?)\s*\}\}")


def _docx_plain_text(docx_path: str) -> str:
    """用 python-docx 拼接正文/表格/页眉页脚的纯文本（自动合并被 Word 拆散的 run）"""
    from docx import Document
    doc = Document(docx_path)
    parts = []

    def walk_paras(paras):
        for p in paras:
            parts.append(p.text)

    def walk_tables(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    walk_paras(cell.paragraphs)
                    walk_tables(cell.tables)

    walk_paras(doc.paragraphs)
    walk_tables(doc.tables)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header,
                   sec.first_page_footer, sec.even_page_header, sec.even_page_footer):
            try:
                walk_paras(hf.paragraphs)
                walk_tables(hf.tables)
            except Exception:
                pass
    return "\n".join(parts)


def extract_placeholders(docx_path: str):
    """提取模板中出现、且系统真实支持自动填充的 {{占位符}}，按出现顺序去重返回"""
    if not docx_path or not os.path.exists(docx_path):
        return []
    try:
        text = _docx_plain_text(docx_path)
    except Exception:
        return []
    found = []
    for m in _PH_RE.finditer(text):
        key = m.group(1).strip()
        if key in SUPPORTED_PLACEHOLDERS and key not in found:
            found.append(key)
    return found


def ensure_builtin_templates():
    """幂等登记内置标准模板；首次（表为空）时设为默认。返回内置模板 id 或 None。"""
    try:
        row = db.query_one("SELECT * FROM templates WHERE name=?", (BUILTIN_NAME,))
        target_path = os.path.join(str(TEMPLATE_DIR), BUILTIN_FILE_NAME)

        if row is None:
            # 复制内置文件到运行时模板目录
            if os.path.exists(BUILTIN_SOURCE):
                shutil.copyfile(BUILTIN_SOURCE, target_path)
            elif not os.path.exists(target_path):
                return None
            empty = db.query_one("SELECT COUNT(*) AS c FROM templates")
            is_default = 1 if (not empty or empty["c"] == 0) else 0
            tid = db.execute(
                "INSERT INTO templates(name, filename, is_default, created_at) VALUES(?,?,?,?)",
                (BUILTIN_NAME, BUILTIN_FILE_NAME, is_default, now_display()))
            return tid

        # 记录存在但磁盘文件丢失：补回
        if not os.path.exists(target_path) and os.path.exists(BUILTIN_SOURCE):
            shutil.copyfile(BUILTIN_SOURCE, target_path)
        return row["id"]
    except Exception:
        return None
