# -*- coding: utf-8 -*-
"""单位管理 + 标签管理 路由"""
import json
import re
from io import BytesIO

import pandas as pd
from flask import Blueprint, request, jsonify, render_template, send_file, Response

import database as db
from core.helpers import content_disposition, now_display
from core.asset_sync import backfill_from_reports

bp = Blueprint("units", __name__)


def _rows(r):
    return dict(r) if r else None


@bp.post("/api/units/sync-from-reports")
def sync_from_reports():
    """一键把存量漏洞报告里的单位 / 域名补录到资产库（幂等，可反复执行）。"""
    try:
        r = backfill_from_reports()
    except Exception as e:
        return jsonify(code=1, msg=f"同步失败：{e}"), 500
    return jsonify(code=0,
                   msg=(f"已扫描 {r['scanned']} 条报告，新增单位 {r['units_created']} 个、"
                        f"域名 {r['domains_created']} 个（已存在的不会重复添加）"),
                   data=r)


# ---------------- 页面 ----------------
@bp.get("/units")
def index():
    return render_template("units.html")


@bp.get("/tags")
def tags_page():
    return render_template("tags.html")


# ---------------- 单位列表 ----------------
@bp.get("/api/units")
def unit_list():
    page = max(1, int(request.args.get("page", 1)))
    limit = min(100, max(1, int(request.args.get("limit", 10))))
    kw = request.args.get("enterprise_name", "").strip()
    tag_id = request.args.get("tag_id", "").strip()

    sql = "SELECT * FROM units WHERE 1=1"
    params = []
    if kw:
        sql += " AND enterprise_name LIKE ?"
        params.append(f"%{kw}%")
    if tag_id:
        sql += " AND id IN (SELECT unit_id FROM unit_tags WHERE tag_id=?)"
        params.append(tag_id)
    sql += " ORDER BY id DESC"

    rows, total = db.paginate(sql, params, page, limit)
    data = []
    for u in rows:
        d = dict(u)
        d["tag_names"] = [t["name"] for t in db.query_all(
            "SELECT t.name FROM tags t JOIN unit_tags ut ON t.id=ut.tag_id WHERE ut.unit_id=?", (u["id"],))]
        data.append(d)
    return jsonify(code=0, msg="ok", count=total, data=data)


# ---------------- 新增（textarea 批量） ----------------
def import_units(unit_names):
    """批量导入单位：去空行、去重、get_or_create"""
    cleaned = [u.strip() for u in unit_names if u.strip()]
    seen, unique = set(), []
    for u in cleaned:
        if u not in seen:
            seen.add(u)
            unique.append(u)
    success, existing = 0, 0
    ids = []
    for name in unique:
        row = db.query_one("SELECT id FROM units WHERE enterprise_name=?", (name,))
        if row:
            existing += 1
            ids.append(row["id"])
        else:
            uid = db.execute(
                "INSERT INTO units(enterprise_name, created_at, updated_at) VALUES(?,?,?)",
                (name, now_display(), now_display()))
            success += 1
            ids.append(uid)
    return {"success_count": success, "existing_count": existing, "ids": ids, "unit_names": unique}


@bp.post("/api/units/add")
def unit_add():
    data = request.get_json(silent=True) or {}
    text = (data.get("enterprise_name") or "").strip()
    if not text:
        return jsonify(code=1, msg="请输入至少一个单位"), 400
    result = import_units(text.splitlines())
    return jsonify(code=0, msg=f"导入完成：新增 {result['success_count']} 个，已存在 {result['existing_count']} 个",
                   data=result)


# ---------------- Excel 模板下载 ----------------
@bp.get("/api/units/template")
def units_template():
    headers = [
        "企业名称", "经营状态", "法定代表人", "注册资本", "实缴资本", "所属省份", "所属城市",
        "所属区县", "统一社会信用代码", "所属行业", "曾用名", "注册地址", "备案信息状态", "ICP官网备案信息数量",
    ]
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "单位信息"
    ws.append(headers)
    for col, w in zip("ABCDEFGHIJKLMN", [28, 10, 12, 14, 14, 10, 10, 10, 22, 22, 22, 40, 14, 14]):
        ws.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(buf, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": content_disposition("【智安盾】单位导入模板.xlsx")})


def _parse_ids(raw):
    out = []
    for x in str(raw or "").split(","):
        x = x.strip()
        if x.isdigit():
            out.append(int(x))
    return out


_UNIT_EXPORT_COLS = [
    ("id", "ID"), ("enterprise_name", "企业名称"), ("business_status", "经营状态"),
    ("legal_representative", "法定代表人"), ("registered_capital", "注册资本"),
    ("paid_in_capital", "实缴资本"), ("province", "所属省份"), ("city", "所属城市"),
    ("district", "所属区县"), ("credit_code", "统一社会信用代码"), ("industry", "所属行业"),
    ("former_names", "曾用名"), ("registered_address", "注册地址"),
    ("filing_status", "备案状态"), ("icp_total", "ICP备案数"), ("domain_count", "域名备案数"),
]
_FILING_CN = {"has_filing": "有备案信息", "no_filing": "无备案信息", "unknown": "未开始查询", "": "未开始查询"}


def _query_units_for_export(ids, kw, tag_id):
    sql = "SELECT * FROM units WHERE 1=1"
    params = []
    if ids:
        sql += " AND id IN (%s)" % ",".join("?" * len(ids))
        params.extend(ids)
    else:
        if kw:
            sql += " AND enterprise_name LIKE ?"
            params.append(f"%{kw}%")
        if tag_id:
            sql += " AND id IN (SELECT unit_id FROM unit_tags WHERE tag_id=?)"
            params.append(tag_id)
    sql += " ORDER BY id DESC"
    return db.query_all(sql, params)


@bp.get("/api/units/export")
def units_export():
    """导出单位名单为 Excel：传 ids= 导出勾选单位；否则按当前搜索/标签条件导出。"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    ids = _parse_ids(request.args.get("ids", ""))
    kw = request.args.get("enterprise_name", "").strip()
    tag_id = request.args.get("tag_id", "").strip()
    rows = _query_units_for_export(ids, kw, tag_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "单位名单"
    headers = [h for _, h in _UNIT_EXPORT_COLS] + ["标签"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
        c.alignment = Alignment(horizontal="center", vertical="center")

    for u in rows:
        tags = "、".join(t["name"] for t in db.query_all(
            "SELECT t.name FROM tags t JOIN unit_tags ut ON t.id=ut.tag_id WHERE ut.unit_id=?",
            (u["id"],)))
        line = []
        keys = set(u.keys())
        for key, _ in _UNIT_EXPORT_COLS:
            v = u[key] if key in keys else ""
            if key == "filing_status":
                v = _FILING_CN.get(v or "", v or "")
            if key in ("icp_total", "domain_count"):
                v = v if (v is not None and int(v) >= 0) else ""
            line.append("" if v is None else v)
        line.append(tags)
        ws.append(line)

    widths = [6, 30, 10, 13, 13, 13, 11, 11, 11, 24, 16, 20, 38, 12, 11, 11, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"【智安盾】单位名单_{len(rows)}个单位.xlsx"
    return Response(
        buf,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": content_disposition(fname)})


@bp.get("/api/units/report-ids")
def units_report_ids():
    """根据勾选的单位 id，返回其名下全部漏洞报告（供按单位批量导出报告）。"""
    ids = _parse_ids(request.args.get("ids", ""))
    if not ids:
        return jsonify(code=1, msg="未选择单位"), 400
    placeholders = ",".join("?" * len(ids))
    units = db.query_all(f"SELECT id, enterprise_name FROM units WHERE id IN ({placeholders})", ids)
    names = [u["enterprise_name"] for u in units]
    if not names:
        return jsonify(code=0, data=[], unit_count=0)
    q = " OR ".join(["unit_name=?"] * len(names))
    rows = db.query_all(
        f"SELECT id, unit_name, system_name, vul_name, severity FROM reports WHERE {q} ORDER BY id",
        names)
    return jsonify(code=0, data=[dict(r) for r in rows], unit_count=len(names))


# ---------------- Excel 批量导入/更新 ----------------
@bp.post("/api/units/update")
def units_update():
    file = request.files.get("file")
    if not file:
        return jsonify(code=1, msg="请上传 Excel 文件"), 400
    # 读到 bytes，便于多次解析（兼容 .xlsx / .xls；第三方导出表常有标题行、列名差异）
    try:
        raw = file.read()
    except Exception as e:
        return jsonify(code=1, msg=f"读取上传文件失败：{e}"), 400

    def _norm_col(c):
        return str(c).strip().replace("　", "").replace(" ", "").replace(":", "").replace("：", "")

    # 列名别名 → 系统标准列名（至少要能识别“单位名称”这一列）
    COL_ALIAS = {
        "企业名称": "企业名称", "单位名称": "企业名称", "公司名称": "企业名称",
        "机构名称": "企业名称", "主体名称": "企业名称", "名称": "企业名称", "单位": "企业名称",
        "法人代表": "法定代表人", "法人": "法定代表人",
        "营业执照号": "统一社会信用代码", "信用代码": "统一社会信用代码",
        "注册地": "注册地址", "地址": "注册地址",
        "ICP备案数": "ICP官网备案信息数量", "备案数": "ICP官网备案信息数量",
    }
    NAME_COLS = {"企业名称", "单位名称", "公司名称", "机构名称", "主体名称", "名称", "单位"}

    header_row = 0
    try:
        probe = pd.read_excel(BytesIO(raw), header=None, nrows=10, dtype=str)
        found = None
        for ri in range(probe.shape[0]):
            for ci in range(probe.shape[1]):
                if _norm_col(probe.iat[ri, ci]) in NAME_COLS:
                    found = ri
                    break
            if found is not None:
                break
        if found is not None:
            header_row = found
    except Exception:
        header_row = 0

    try:
        df = pd.read_excel(BytesIO(raw), header=header_row)
    except Exception as e:
        return jsonify(code=1, msg=f"读取 Excel 失败：{e}（请另存为 .xlsx 后使用系统下载的模板导入）"), 400

    # 规范化 + 别名映射列名
    df.columns = [_norm_col(c) for c in df.columns]
    rename_map = {}
    for c in df.columns:
        std = COL_ALIAS.get(c)
        if std and std not in df.columns:
            rename_map[c] = std
    if rename_map:
        df = df.rename(columns=rename_map)

    if "企业名称" not in df.columns:
        return jsonify(code=1, msg="未识别到“企业名称/单位名称”列：请确认表头中含该列（可直接使用系统下载的导入模板）"), 400

    field_map = {
        "企业名称": "enterprise_name", "经营状态": "business_status", "法定代表人": "legal_representative",
        "注册资本": "registered_capital", "实缴资本": "paid_in_capital", "所属省份": "province",
        "所属城市": "city", "所属区县": "district", "统一社会信用代码": "credit_code",
        "所属行业": "industry", "曾用名": "former_names", "注册地址": "registered_address",
        "备案信息状态": "filing_status", "ICP官网备案信息数量": "icp_total",
    }
    invalid_status = ["注销", "关闭", "废止", "吊销", "迁出", "停业", "歇业", "停业整顿", "经营异常"]
    create_count = update_count = skip_count = 0
    now = now_display()

    for _, row in df.iterrows():
        name = str(row.get("企业名称", "")).strip()
        if not name or name.lower() == "nan":
            continue

        existing = db.query_one("SELECT id FROM units WHERE enterprise_name=?", (name,))
        uid = None
        if existing:
            uid = existing["id"]
        else:
            # 企业名称匹配不到时，按“曾用名”列（分号分隔）在 企业名称/曾用名 字段中做边界正则匹配
            former_raw = row.get("曾用名", "")
            former_names = []
            if pd.notna(former_raw):
                former_names = [x.strip() for x in str(former_raw).replace(",", ";").split(";") if x.strip()]
            for fn in former_names:
                pat = r"(^|;)\s*" + re.escape(fn) + r"\s*($|;)"
                cand = db.query_one(
                    "SELECT id, enterprise_name, former_names FROM units "
                    "WHERE enterprise_name LIKE ? OR former_names LIKE ? "
                    "ORDER BY id LIMIT 50", (f"%{fn}%", f"%{fn}%"))
                hit = None
                if cand:
                    for c in db.query_all(
                            "SELECT id, enterprise_name, former_names FROM units "
                            "WHERE enterprise_name LIKE ? OR former_names LIKE ?",
                            (f"%{fn}%", f"%{fn}%")):
                        if re.search(pat, c["enterprise_name"] or "") or re.search(pat, c["former_names"] or ""):
                            hit = c
                            break
                if hit:
                    # 曾用名匹配成功：把企业名称更新为最新名称，并保留旧名到曾用名
                    old_name = hit["enterprise_name"] or ""
                    old_former = hit["former_names"] or ""
                    merged = ";".join([x for x in [old_name, old_former] if x and x != name])
                    db.execute("UPDATE units SET enterprise_name=?, former_names=?, updated_at=? WHERE id=?",
                               (name, merged, now, hit["id"]))
                    uid = hit["id"]
                    break

        if uid is None:
            # 仍未匹配：按经营状态决定是否新增（注销/关闭等异常状态跳过）
            status = row.get("经营状态", "") if "经营状态" in df.columns else ""
            if pd.isna(status):
                status = ""
            status = str(status).strip()
            # 只有“明确标注”为注销/吊销/关闭等异常状态的新单位才跳过；
            # 经营状态留空（使用仅含企业名称的简易模板）视为状态未知，应正常新增。
            if status and any(bad in status for bad in invalid_status):
                skip_count += 1
                continue
            uid = db.execute("INSERT INTO units(enterprise_name, created_at, updated_at) VALUES(?,?,?)",
                             (name, now, now))
            create_count += 1

        updates, params = [], []
        for cn, field in field_map.items():
            if cn == "企业名称":
                continue
            if cn in df.columns:
                val = row.get(cn)
                if pd.notna(val) and val != "":
                    if cn == "ICP官网备案信息数量":
                        try:
                            val = int(float(val))
                        except Exception:
                            continue
                    updates.append(f"{field}=?")
                    params.append(str(val))
        if updates:
            db.execute(f"UPDATE units SET {','.join(updates)}, updated_at=? WHERE id=?", (*params, now, uid))
            update_count += 1

    return jsonify(code=0, msg=f"处理完成：新增 {create_count} 条，更新 {update_count} 条，跳过 {skip_count} 条")


# ---------------- 单位详情/更新/删除 ----------------
@bp.get("/api/unit/detail")
def unit_detail():
    uid = request.args.get("id")
    row = db.query_one("SELECT * FROM units WHERE id=?", (uid,))
    if not row:
        return jsonify(code=1, msg="单位不存在"), 404
    return jsonify(code=0, data=_rows(row))


@bp.post("/api/unit/update")
def unit_update():
    data = request.get_json(silent=True) or {}
    uid = data.get("id")
    if not uid:
        return jsonify(code=1, msg="缺少 id"), 400
    updatable = ["enterprise_name", "business_status", "legal_representative", "registered_capital",
                 "paid_in_capital", "province", "city", "district", "credit_code", "industry",
                 "former_names", "registered_address", "filing_status", "icp_total", "domain_count"]
    sets, params = [], []
    for f in updatable:
        if f in data and data[f] is not None:
            sets.append(f"{f}=?")
            params.append(data[f])
    if not sets:
        return jsonify(code=1, msg="没有可更新的字段")
    params.append(now_display())
    params.append(uid)
    db.execute(f"UPDATE units SET {','.join(sets)}, updated_at=? WHERE id=?", (*params,))
    return jsonify(code=0, msg="更新成功")


@bp.post("/api/unit/delete")
def unit_delete():
    data = request.get_json(silent=True) or {}
    uid = data.get("id")
    if not uid:
        return jsonify(code=1, msg="缺少 id"), 400
    db.execute("DELETE FROM unit_tags WHERE unit_id=?", (uid,))
    db.execute("DELETE FROM units WHERE id=?", (uid,))
    return jsonify(code=0, msg="删除成功")


# ---------------- ICP 数据包导入 ----------------
def _is_valid_company_name(name):
    if re.match(r"^(?:\d{1,3}\.){3}\d{1,3}$", name):
        return False
    if re.match(r"^(?:[a-zA-Z0-9-]+\.)+[a-zA-Z]{2,}$", name):
        return False
    return True


@bp.post("/api/units/import-icp")
def unit_import_icp():
    data = request.get_json(silent=True) or {}
    text = (data.get("data") or "").strip()
    if not text:
        return jsonify(code=1, msg="缺少数据"), 400
    from routes.domains_bp import import_domains_from_json

    result = {}
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        company, payload = None, None
        # 新格式：整行 JSON 含 query_keyword
        try:
            j = json.loads(line)
            if isinstance(j, dict) and "query_keyword" in j and "result" in j:
                company = j["query_keyword"]
                payload = j["result"]
        except json.JSONDecodeError:
            # 旧格式：公司名: {json}
            if ":" in line:
                company = line.split(":", 1)[0].strip()
                try:
                    payload = json.loads(line.split(":", 1)[1].strip())
                except json.JSONDecodeError:
                    payload = None
        if not company or not payload:
            continue
        if not _is_valid_company_name(company):
            result[company] = "无效的公司名称（域名或IP地址）"
            continue
        params = payload.get("params", {}) if isinstance(payload, dict) else {}
        total = params.get("total", 0)
        # 先确认单位存在（db.execute 返回 lastrowid，不能用于判断 UPDATE 命中行数）
        hit = db.query_one("SELECT id FROM units WHERE enterprise_name=?", (company,))
        if not hit:
            result[company] = "单位不存在，未更新"
            continue
        db.execute("UPDATE units SET icp_total=? WHERE id=?", (total, hit["id"]))
        domain_list = params.get("list", [])
        if domain_list:
            import_domains_from_json(domain_list)
        result[company] = total
    return jsonify(result)


# ---------------- 批量查询单位 ----------------
@bp.post("/api/unit/batch-query")
def unit_batch_query():
    data = request.get_json(silent=True) or {}
    values = data.get("values", [])
    if not values:
        return jsonify(code=1, msg="缺少查询值")
    # 支持多值（OR 模糊匹配）
    sql = "SELECT * FROM units WHERE " + " OR ".join(["enterprise_name LIKE ?"] * len(values))
    rows = db.query_all(sql, [f"%{v}%" for v in values])
    return jsonify(code=0, data=[dict(r) for r in rows])


# ---------------- 标签管理 ----------------
@bp.get("/api/tags")
def tag_list():
    page = max(1, int(request.args.get("page", 1)))
    limit = min(100, max(1, int(request.args.get("limit", 10))))
    kw = request.args.get("keyword", "").strip()
    sql = "SELECT * FROM tags WHERE 1=1"
    params = []
    if kw:
        sql += " AND name LIKE ?"
        params.append(f"%{kw}%")
    sql += " ORDER BY id"
    rows, total = db.paginate(sql, params, page, limit)
    data = []
    for t in rows:
        units = db.query_all(
            "SELECT u.enterprise_name FROM units u JOIN unit_tags ut ON u.id=ut.unit_id WHERE ut.tag_id=? ORDER BY u.id",
            (t["id"],))
        data.append({**dict(t), "unit": "\n".join(u["enterprise_name"] for u in units),
                     "unit_count": len(units)})
    return jsonify(code=0, msg="ok", count=total, data=data)


@bp.get("/api/tags/options")
def tag_options():
    rows = db.query_all("SELECT id, name FROM tags ORDER BY id")
    return jsonify(code=0, data=[dict(r) for r in rows])


@bp.post("/api/tags/add")
def tag_add():
    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify(code=1, msg="标签名称不能为空")
    if db.query_one("SELECT id FROM tags WHERE name=?", (name,)):
        return jsonify(code=1, msg="标签已存在")
    tid = db.execute("INSERT INTO tags(name, description, created_at) VALUES(?,?,?)",
                     (name, data.get("description", ""), now_display()))
    unit_names = [u.strip() for u in (data.get("unit") or "").splitlines() if u.strip()]
    for un in unit_names:
        row = db.query_one("SELECT id FROM units WHERE enterprise_name=?", (un,))
        if row:
            db.execute("INSERT OR IGNORE INTO unit_tags(unit_id, tag_id) VALUES(?,?)", (row["id"], tid))
    return jsonify(code=0, msg="添加成功", tag_id=tid)


@bp.get("/api/tags/detail")
def tag_detail():
    tid = request.args.get("id")
    t = db.query_one("SELECT * FROM tags WHERE id=?", (tid,))
    if not t:
        return jsonify(code=1, msg="标签不存在")
    units = db.query_all(
        "SELECT u.enterprise_name FROM units u JOIN unit_tags ut ON u.id=ut.unit_id WHERE ut.tag_id=?",
        (tid,))
    return jsonify(code=0, data={**dict(t), "unit": "\n".join(u["enterprise_name"] for u in units)})


@bp.post("/api/tags/update")
def tag_update():
    data = request.get_json(silent=True) or {}
    tid = data.get("id")
    if not tid:
        return jsonify(code=1, msg="缺少 id")
    if not db.query_one("SELECT id FROM tags WHERE id=?", (tid,)):
        return jsonify(code=1, msg="标签不存在")
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify(code=1, msg="标签名称不能为空")
    db.execute("UPDATE tags SET name=?, description=? WHERE id=?", (name, data.get("description", ""), tid))
    # 与原工具一致：重建关联，单位不存在则自动创建
    db.execute("DELETE FROM unit_tags WHERE tag_id=?", (tid,))
    unit_names = [u.strip() for u in (data.get("unit") or "").splitlines() if u.strip()]
    imp = import_units(unit_names)
    for uid in imp["ids"]:
        db.execute("INSERT OR IGNORE INTO unit_tags(unit_id, tag_id) VALUES(?,?)", (uid, tid))
    return jsonify(code=0, msg="更新成功")


@bp.post("/api/tags/assign")
def tag_assign():
    """给指定标签追加关联单位（不清空已有关联）；单位不存在时自动创建。对齐原工具 tag_assign。"""
    data = request.get_json(silent=True) or {}
    tid = data.get("tag_id")
    unit_text = (data.get("unit") or "").strip()
    if not tid or not unit_text:
        return jsonify(code=1, msg="标签和单位不能为空")
    if not db.query_one("SELECT id FROM tags WHERE id=?", (tid,)):
        return jsonify(code=1, msg="标签不存在")
    unit_names = [u.strip() for u in unit_text.splitlines() if u.strip()]
    imp = import_units(unit_names)
    for uid in imp["ids"]:
        db.execute("INSERT OR IGNORE INTO unit_tags(unit_id, tag_id) VALUES(?,?)", (uid, tid))
    return jsonify(code=0, msg="关联成功", success_count=imp["success_count"],
                   existing_count=imp["existing_count"], units=imp["unit_names"])


@bp.get("/api/unit/tags")
def unit_tags_get():
    """返回某单位当前已关联的标签 id 列表（供单位列表打标签回显）。"""
    uid = request.args.get("unit_id")
    if not uid or not db.query_one("SELECT id FROM units WHERE id=?", (uid,)):
        return jsonify(code=1, msg="单位不存在")
    ids = [r["tag_id"] for r in db.query_all(
        "SELECT tag_id FROM unit_tags WHERE unit_id=?", (uid,))]
    return jsonify(code=0, data={"tag_ids": ids})


@bp.post("/api/unit/set-tags")
def unit_set_tags():
    """覆盖式设置某单位的标签集合：勾选即关联、取消即解除，单位列表与标签管理实时同步。"""
    data = request.get_json(silent=True) or {}
    uid = data.get("unit_id")
    tag_ids = data.get("tag_ids") or []
    if not uid or not db.query_one("SELECT id FROM units WHERE id=?", (uid,)):
        return jsonify(code=1, msg="单位不存在")
    clean = []
    for t in tag_ids:
        try:
            clean.append(int(t))
        except Exception:
            continue
    db.execute("DELETE FROM unit_tags WHERE unit_id=?", (uid,))
    for tid in dict.fromkeys(clean):
        if db.query_one("SELECT id FROM tags WHERE id=?", (tid,)):
            db.execute("INSERT OR IGNORE INTO unit_tags(unit_id, tag_id) VALUES(?,?)", (uid, tid))
    return jsonify(code=0, msg="标签已保存")


@bp.post("/api/tags/delete")
def tag_delete():
    data = request.get_json(silent=True) or {}
    tid = data.get("id")
    db.execute("DELETE FROM unit_tags WHERE tag_id=?", (tid,))
    db.execute("DELETE FROM tags WHERE id=?", (tid,))
    return jsonify(code=0, msg="删除成功")
