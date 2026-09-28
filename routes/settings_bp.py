# -*- coding: utf-8 -*-
"""系统设置路由（外观 / AI 配置 / 数据管理 / 内置模板）"""
import os
import uuid

from flask import Blueprint, request, jsonify, render_template, send_file

import database as db
from config import get_cfg, save_config, BG_DIR, BASE_DIR
from core.helpers import now_display, content_disposition

bp = Blueprint("settings", __name__)


@bp.get("/settings")
def page():
    return render_template("settings.html")


# ---------------- 外观设置 ----------------
@bp.get("/api/settings/appearance")
def get_appearance():
    cfg = get_cfg()
    bgs = []
    for f in sorted(os.listdir(str(BG_DIR))):
        if f.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")):
            bgs.append(f)
    return jsonify(code=0, data={
        "app_title": cfg.get("app_title", ""),
        "app_subtitle": cfg.get("app_subtitle", ""),
        "theme_color": cfg.get("theme_color", "#2563eb"),
        "background": cfg.get("background", ""),
        "background_overlay": cfg.get("background_overlay", 0.35),
        "background_blur": cfg.get("background_blur", 0),
        "report_header": cfg.get("report_header", "网络安全测试报告"),
        "export_org": cfg.get("export_org", ""),
        "backgrounds": bgs,
    })


@bp.post("/api/settings/appearance")
def save_appearance():
    data = request.get_json(silent=True) or {}
    cfg = get_cfg()
    for k in ("app_title", "app_subtitle", "theme_color", "background",
              "report_header", "export_org"):
        if k in data:
            cfg[k] = data[k]
    try:
        cfg["background_overlay"] = min(0.9, max(0, float(data.get("background_overlay", 0.35))))
        cfg["background_blur"] = min(30, max(0, int(data.get("background_blur", 0))))
    except Exception:
        pass
    save_config(cfg)
    return jsonify(code=0, msg="外观设置已保存")


@bp.post("/api/settings/background/upload")
def upload_background():
    f = request.files.get("file")
    if not f:
        return jsonify(code=1, msg="未选择文件")
    ext = os.path.splitext(f.filename or "")[1].lower()
    if ext not in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
        return jsonify(code=1, msg="不支持的图片格式")
    filename = f"bg_{uuid.uuid4().hex}{ext}"
    f.save(os.path.join(str(BG_DIR), filename))
    cfg = get_cfg()
    cfg["background"] = filename
    save_config(cfg)
    return jsonify(code=0, msg="背景图已上传并应用", filename=filename)


@bp.post("/api/settings/background/delete")
def delete_background():
    data = request.get_json(silent=True) or {}
    name = data.get("name", "")
    if not name or ".." in name or "/" in name or "\\" in name:
        return jsonify(code=1, msg="非法文件名")
    path = os.path.join(str(BG_DIR), name)
    if os.path.exists(path):
        os.remove(path)
    cfg = get_cfg()
    if cfg.get("background") == name:
        cfg["background"] = ""
        save_config(cfg)
    return jsonify(code=0, msg="已删除")


@bp.get("/media/backgrounds/<path:filename>")
def bg_file(filename):
    safe = os.path.basename(filename)
    path = os.path.join(str(BG_DIR), safe)
    if not os.path.exists(path):
        return jsonify(code=1, msg="文件不存在"), 404
    return send_file(path)


# ---------------- AI 配置 ----------------
@bp.get("/api/settings/ai")
def get_ai():
    ai = get_cfg().get("ai", {})
    key = ai.get("api_key", "")
    return jsonify(code=0, data={
        "base_url": ai.get("base_url", ""),
        "model": ai.get("model", ""),
        "api_key": key,  # 回显明文（本地工具，便于修改）
        "proxy": ai.get("proxy", ""),
        "timeout": ai.get("timeout", 120),
    })


@bp.post("/api/settings/ai")
def save_ai():
    data = request.get_json(silent=True) or {}
    cfg = get_cfg()
    ai = cfg.setdefault("ai", {})
    if "base_url" in data:
        ai["base_url"] = data["base_url"]
    if "model" in data:
        ai["model"] = data["model"]
    if "proxy" in data:
        ai["proxy"] = (data["proxy"] or "").strip()
    if "api_key" in data:
        ai["api_key"] = data["api_key"].strip()
    if "timeout" in data:
        try:
            ai["timeout"] = int(data["timeout"])
        except Exception:
            pass
    save_config(cfg)
    # AI 配置是全系统唯一配置源：保存后自动同步给 Shelling（best-effort，失败不影响保存）
    synced, sync_msg = _sync_ai_to_shelling(cfg.get("ai", {}))
    msg = "AI 配置已保存"
    if synced is True:
        msg += "；已同步到 Shelling（" + sync_msg + "）"
    elif synced is False:
        msg += "；同步到 Shelling 失败：" + sync_msg
    elif sync_msg:
        msg += "；" + sync_msg
    return jsonify(code=0, msg=msg, shelling_synced=synced)


# ---------------- 统一 AI 配置：同步到 Shelling ----------------
def _sync_ai_to_shelling(ai: dict):
    """
    把智安盾的 AI 配置推送到 Shelling 的 LLM 配置（幂等）。
    返回 (是否成功, 说明)；api_key 为空时返回 (None, 说明) 表示跳过。
    """
    from core import shelling_client
    base_url = (ai.get("base_url") or "").strip()
    model = (ai.get("model") or "").strip()
    api_key = (ai.get("api_key") or "").strip()
    if not api_key:
        # 没有 Key 时若强行同步，Shelling 会拿着空 Key 去调用并失败，比「不配置」更糟
        return None, "未填写 API Key，已跳过同步"
    if not base_url or not model:
        return None, "Base URL 或模型为空，已跳过同步"
    try:
        res = shelling_client.get_client().sync_llm_config(base_url, api_key, model)
    except shelling_client.ShellingError as e:
        return False, str(e)
    except Exception as e:
        return False, f"同步异常：{e}"
    cfg = res.get("config") or {}
    return True, f"{'已创建' if res.get('action') == 'created' else '已更新'}配置「{cfg.get('name')}」（模型 {cfg.get('model')}）"


@bp.get("/api/settings/ai/shelling")
def shelling_llm_status():
    """查看 Shelling 侧的 LLM 配置状态（是否已与智安盾统一）"""
    from core import shelling_client
    try:
        st = shelling_client.get_client().llm_status()
    except Exception as e:
        st = {"reachable": False, "error": f"{e}", "total": 0,
              "synced": None, "active_main": None, "active_sub": None}
    return jsonify(code=0, data=st)


@bp.post("/api/settings/ai/sync")
def sync_ai_to_shelling():
    """手动把智安盾的 AI 配置同步到 Shelling，并让 Shelling 实测一次"""
    from core import shelling_client
    data = request.get_json(silent=True) or {}
    ai = get_cfg().get("ai", {})
    # 表单里尚未保存的值也允许直接同步
    base_url = (data.get("base_url") or ai.get("base_url") or "").strip()
    model = (data.get("model") or ai.get("model") or "").strip()
    api_key = (data.get("api_key") or ai.get("api_key") or "").strip()
    if not api_key:
        return jsonify(code=1, msg="请先填写 AI API Key —— Shelling 侧需要一个可用的 Key 才能真正启用 AI 分析")
    try:
        res = shelling_client.get_client().sync_llm_config(base_url, api_key, model)
    except shelling_client.ShellingError as e:
        return jsonify(code=1, msg=str(e))
    except Exception as e:
        return jsonify(code=1, msg=f"同步失败：{e}")

    cfg = res.get("config") or {}
    took_over = res.get("took_over") or []
    msg = f"{'已创建' if res.get('action') == 'created' else '已更新'} Shelling 侧配置「{cfg.get('name')}」（模型 {cfg.get('model')}）"
    if took_over:
        msg += "，已接管原有生效配置：" + "、".join(took_over)

    verify = data.get("verify", True)
    test = None
    if verify and cfg.get("id"):
        try:
            test = shelling_client.get_client().test_llm_config(cfg["id"])
        except shelling_client.ShellingError as e:
            test = {"success": False, "message": "验证请求失败", "error": str(e)}
        if test.get("success"):
            msg += "；Shelling 侧实测调用成功"
        else:
            msg += "；但 Shelling 侧实测失败：" + str(test.get("error") or test.get("message") or "")[:160]

    return jsonify(code=0, msg=msg, data={"action": res.get("action"), "config": cfg,
                                          "took_over": took_over, "test": test})


# ---------------- Shelling 扫描平台配置 ----------------
@bp.get("/api/settings/shelling")
def get_shelling():
    sh = get_cfg().get("shelling", {})
    return jsonify(code=0, data={
        "base_url": sh.get("base_url", ""),
        "username": sh.get("username", ""),
        "password": sh.get("password", ""),  # 回显明文（本地工具，便于修改）
        "timeout": sh.get("timeout", 600),
    })


@bp.post("/api/settings/shelling")
def save_shelling():
    data = request.get_json(silent=True) or {}
    cfg = get_cfg()
    sh = cfg.setdefault("shelling", {})
    for k in ("base_url", "username"):
        if k in data:
            sh[k] = (data[k] or "").strip()
    if "password" in data:
        sh["password"] = (data["password"] or "").strip()
    if "timeout" in data:
        try:
            sh["timeout"] = max(30, int(data["timeout"]))
        except Exception:
            pass
    save_config(cfg)
    # 配置变更后丢弃缓存的客户端与 JWT
    from core import shelling_client
    shelling_client.reset_client()
    return jsonify(code=0, msg="Shelling 配置已保存")


@bp.post("/api/settings/shelling/test")
def test_shelling():
    """测试 Shelling 平台连通性：优先用表单当前值，未填则回落到已保存配置"""
    from core import shelling_client
    data = request.get_json(silent=True) or {}
    try:
        client = shelling_client.make_client(
            base_url=(data.get("base_url") or "").strip() or None,
            username=(data.get("username") or "").strip() or None,
            password=(data.get("password") or "").strip() or None,
            timeout=data.get("timeout"),
        )
        info = client.ping()
    except shelling_client.ShellingError as e:
        return jsonify(code=1, msg=str(e))
    except Exception as e:
        return jsonify(code=1, msg=f"连接失败：{e}")
    return jsonify(code=0, msg="连接正常", data={
        "base_url": info["base_url"],
        "scanner_count": info["scanner_count"],
        "scanners": info["scanners"][:8],
    })


# ---------------- 数据管理 ----------------
@bp.post("/api/settings/seed-demo")
def seed_demo():
    return jsonify(**db.seed_demo_data())


@bp.post("/api/settings/clear-data")
def clear_data():
    return jsonify(**db.clear_all_data())


def _bring_folder_to_front(folder_basename: str):
    """Windows：找到标题为该文件夹名的资源管理器窗口并强制置前。
    后台 Flask 进程直接 SetForegroundWindow 通常只闪任务栏，附加到当前前台线程后可成功前置。
    explorer 新开窗口是异步的，需短暂轮询等窗口真正出现。"""
    import time
    try:
        import ctypes
        from ctypes import wintypes
        u = ctypes.windll.user32
        k32 = ctypes.windll.kernel32
        u.GetWindowTextW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
        u.GetClassNameW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
        WNDENUMPROC = ctypes.WINFUNCTYPE(ctypes.c_bool, wintypes.HWND, wintypes.LPARAM)

        def find_hwnd():
            found = []

            def _cb(hwnd, _l):
                if u.IsWindowVisible(hwnd):
                    tb, cb_ = ctypes.create_unicode_buffer(512), ctypes.create_unicode_buffer(256)
                    u.GetWindowTextW(hwnd, tb, 512)
                    u.GetClassNameW(hwnd, cb_, 256)
                    title, cls = tb.value, cb_.value
                    head = title.split(" - ")[0].strip() if title else ""
                    if cls in ("CabinetWClass", "ExploreWClass") and title and head == folder_basename:
                        found.append(hwnd)
                return True

            u.EnumWindows(WNDENUMPROC(_cb), 0)
            return found[0] if found else None

        # explorer 新开窗口有延迟，轮询最多约 1.5s 等它出现
        hwnd = None
        for _ in range(15):
            hwnd = find_hwnd()
            if hwnd:
                break
            time.sleep(0.1)
        if not hwnd:
            return
        u.ShowWindow(hwnd, 9)  # SW_RESTORE
        fg = u.GetForegroundWindow()
        cur_tid = k32.GetCurrentThreadId()
        fg_tid = u.GetWindowThreadProcessId(fg, None) if fg else 0
        attached = False
        if fg_tid and fg_tid != cur_tid:
            attached = bool(u.AttachThreadInput(cur_tid, fg_tid, True))
        try:
            u.ShowWindow(hwnd, 5)  # SW_SHOW
            u.BringWindowToTop(hwnd)
            u.SetForegroundWindow(hwnd)
        finally:
            if attached:
                u.AttachThreadInput(cur_tid, fg_tid, False)
    except Exception:
        # 前置仅为体验增强，失败不影响“窗口已打开”这一事实
        pass


@bp.get("/api/settings/data-path")
def get_data_path():
    from config import DATA_DIR
    return jsonify(code=0, path=str(DATA_DIR), exists=os.path.isdir(str(DATA_DIR)))


@bp.post("/api/settings/open-data-dir")
def open_data_dir():
    """在本机文件管理器中打开数据目录（本地工具专用，网页本身无权访问磁盘）。"""
    import sys
    import subprocess
    from config import DATA_DIR
    target = str(DATA_DIR)
    errors = []
    try:
        os.makedirs(target, exist_ok=True)
        if os.name == "nt":
            # 后台无窗口进程下 os.startfile 常把文件夹开在浏览器后面；
            # 优先显式启动 explorer（会激活已开窗口），再用 Win32 强制置前；逐级兜底。
            try:
                subprocess.Popen(["explorer.exe", target],
                                 stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            except Exception as e:
                errors.append(f"explorer: {e}")
                try:
                    os.startfile(target)  # noqa
                except Exception as e2:
                    errors.append(f"startfile: {e2}")
            _bring_folder_to_front(os.path.basename(target.rstrip("\\/")))
        elif sys.platform == "darwin":
            subprocess.Popen(["open", target])
        else:
            subprocess.Popen(["xdg-open", target])
    except Exception as e:
        return jsonify(code=1, msg=f"打开数据目录失败：{e}", path=target), 500
    if os.name == "nt" and len(errors) == 2:
        return jsonify(code=1, msg="打开数据目录失败：" + "；".join(errors), path=target), 500
    return jsonify(code=0, msg="已打开数据目录（若窗口未到最前，请查看任务栏的文件夹图标）", path=target)


# ---------------- 系统信息 ----------------
@bp.get("/api/settings/stats")
def stats():
    conn = db.get_db()
    cur = conn.cursor()
    n_units = cur.execute("SELECT COUNT(*) FROM units").fetchone()[0]
    n_domains = cur.execute("SELECT COUNT(*) FROM domains WHERE is_deleted=0").fetchone()[0]
    n_vulns = cur.execute("SELECT COUNT(*) FROM vulnerabilities").fetchone()[0]
    n_reports = cur.execute("SELECT COUNT(*) FROM reports WHERE status=0").fetchone()[0]
    n_templates = cur.execute("SELECT COUNT(*) FROM templates").fetchone()[0]
    severity = dict(cur.execute(
        "SELECT severity, COUNT(*) FROM reports WHERE status=0 GROUP BY severity").fetchall())
    cur.close()
    return jsonify(code=0, data={
        "units": n_units, "domains": n_domains, "vulns": n_vulns,
        "reports": n_reports, "templates": n_templates, "severity": severity,
    })


# ---------------- 内置模板下载（普通/公文） ----------------
@bp.get("/api/settings/builtin-template/<key>")
def builtin_template(key):
    """内置模板下载：guide/standard 直接下发随系统分发的标准占位符模板（可直接编辑后上传）"""
    if key in ("guide", "standard"):
        from config import BUILTIN_TEMPLATE_DIR
        path = os.path.join(str(BUILTIN_TEMPLATE_DIR), "标准漏洞报告模板.docx")
        if os.path.exists(path):
            return send_file(
                path, as_attachment=True,
                download_name="智安盾-标准漏洞报告模板（含占位符）.docx",
                mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    """导出内置模板的 docx 说明文件，便于用户参考占位符"""
    import io
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph("智安盾 · 自定义报告模板占位符说明")
    for run in p.runs:
        run.font.size = Pt(18)
        run.font.bold = True
    doc.add_paragraph("将本文件中的占位符替换为你自己的模板内容后另存为 .docx 上传即可。")
    doc.add_paragraph("")
    doc.add_paragraph("支持的占位符：")
    for ph in ["单位名称", "漏洞名称", "系统名称", "域名", "漏洞地址", "漏洞描述",
               "检测过程", "漏洞等级", "漏洞危害", "整改建议", "归属证明", "报告日期", "单位落款"]:
        doc.add_paragraph(f"{{{{{ph}}}}}")
    doc.add_paragraph("")
    doc.add_paragraph("示例模板骨架：")
    doc.add_paragraph("一、基本信息")
    doc.add_paragraph("单位名称：{{单位名称}}")
    doc.add_paragraph("系统名称：{{系统名称}}")
    doc.add_paragraph("二、漏洞详情")
    doc.add_paragraph("目标：{{漏洞地址}}")
    doc.add_paragraph("漏洞描述：{{漏洞描述}}")
    doc.add_paragraph("检测过程：{{检测过程}}")
    doc.add_paragraph("漏洞等级：{{漏洞等级}}")
    doc.add_paragraph("三、整改建议")
    doc.add_paragraph("{{整改建议}}")
    doc.add_paragraph("")
    doc.add_paragraph("报告日期：{{报告日期}}")
    doc.add_paragraph("落款：{{单位落款}}")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name="智安盾-自定义模板占位符说明.docx",
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
